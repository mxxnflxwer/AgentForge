import io
from docx import Document as DocxDocument
from pypdf import PdfWriter
import pytest
from fastapi.testclient import TestClient

from app.models.document import DocumentProcessingStatus
from app.services.document.pipeline import get_document_pipeline


def register_and_login_user(client: TestClient, email: str, name: str = "Test User") -> str:
    """Helper to register and log in a user, returning the session ID cookie."""
    client.post(
        "/api/auth/register",
        json={
            "email": email,
            "password": "StrongPassword123!",
            "name": name,
        },
    )
    res = client.post(
        "/api/auth/login",
        json={
            "email": email,
            "password": "StrongPassword123!",
        },
    )
    assert res.status_code == 200
    session_cookie = res.cookies.get("session_id")
    return session_cookie


def create_sample_docx_bytes() -> bytes:
    doc = DocxDocument()
    doc.add_heading("MEDICAL REPORT", level=1)
    doc.add_paragraph("CHIEF COMPLAINT: Recurrent headaches.")
    doc.add_paragraph("ASSESSMENT & PLAN: Prescribe preventative medication.")
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def create_sample_pdf_bytes() -> bytes:
    content = (
        "%PDF-1.4\n"
        "1 0 obj\n"
        "<< /Type /Catalog /Pages 2 0 R >>\n"
        "endobj\n"
        "2 0 obj\n"
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>\n"
        "endobj\n"
        "3 0 obj\n"
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>\n"
        "endobj\n"
        "4 0 obj\n"
        "<< /Length 55 >>\n"
        "stream\n"
        "BT\n"
        "/F1 12 Tf\n"
        "100 700 Td\n"
        "(CHIEF COMPLAINT: Chest Pain)\n"
        "Tj\n"
        "ET\n"
        "endstream\n"
        "endobj\n"
        "5 0 obj\n"
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>\n"
        "endobj\n"
        "xref\n"
        "0 6\n"
        "0000000000 65535 f \n"
        "0000000009 00000 n \n"
        "0000000058 00000 n \n"
        "0000000115 00000 n \n"
        "0000000234 00000 n \n"
        "0000000338 00000 n \n"
        "trailer\n"
        "<< /Size 6 /Root 1 0 R >>\n"
        "startxref\n"
        "427\n"
        "%%EOF"
    )
    return content.encode("latin1")




def test_unauthenticated_document_access(client: TestClient):
    res = client.get("/api/documents")
    assert res.status_code == 401
    assert "Authentication required" in res.json()["detail"]

    res_post = client.post(
        "/api/documents",
        files={"file": ("test.txt", b"Some text content", "text/plain")},
    )
    assert res_post.status_code == 401


def test_document_upload_validation(client: TestClient, db_session):
    cookie_a = register_and_login_user(client, "uploader@example.com", "Uploader")
    client.cookies.set("session_id", cookie_a)

    # 1. Empty file
    res = client.post(
        "/api/documents",
        files={"file": ("empty.txt", b"", "text/plain")},
    )
    assert res.status_code == 400
    assert "empty file" in res.json()["detail"].lower()

    # 2. Unsupported extension
    res = client.post(
        "/api/documents",
        files={"file": ("malicious.exe", b"MZbinarycontent", "application/octet-stream")},
    )
    assert res.status_code == 400
    assert "unsupported file extension" in res.json()["detail"].lower()

    # 3. Invalid magic header for PDF
    res = client.post(
        "/api/documents",
        files={"file": ("fake.pdf", b"Not a real pdf content", "application/pdf")},
    )
    assert res.status_code == 400
    assert "do not match valid PDF format" in res.json()["detail"]


def test_upload_and_process_txt_document(client: TestClient, db_session):
    cookie_a = register_and_login_user(client, "txt_user@example.com", "TXT User")
    client.cookies.set("session_id", cookie_a)

    content = (
        "CHIEF COMPLAINT:\n"
        "Patient presents with fever and cough.\n\n"
        "HISTORY OF PRESENT ILLNESS:\n"
        "Started 3 days ago with mild symptoms.\n\n"
        "ASSESSMENT & PLAN:\n"
        "Prescribe rest and hydration."
    ).encode("utf-8")

    res = client.post(
        "/api/documents",
        data={"specialty": "Pulmonology"},
        files={"file": ("clinical_note.txt", content, "text/plain")},
    )
    assert res.status_code == 201
    doc_data = res.json()
    assert doc_data["filename"] == "clinical_note.txt"
    assert doc_data["file_type"] == "txt"
    assert doc_data["specialty"] == "Pulmonology"
    doc_id = doc_data["id"]

    # Trigger synchronous pipeline to verify processing state
    pipeline = get_document_pipeline()
    pipeline.process_document(db=db_session, document_id=doc_id)

    # Fetch document detail
    detail_res = client.get(f"/api/documents/{doc_id}")
    assert detail_res.status_code == 200
    detail = detail_res.json()
    assert detail["status"] == DocumentProcessingStatus.COMPLETED.value
    assert detail["chunk_count"] >= 3
    assert len(detail["chunks"]) == detail["chunk_count"]
    assert any("CHIEF COMPLAINT" in c["section_title"].upper() for c in detail["chunks"] if c["section_title"])


def test_upload_and_process_docx_document(client: TestClient, db_session):
    cookie = register_and_login_user(client, "docx_user@example.com", "DOCX User")
    client.cookies.set("session_id", cookie)

    docx_bytes = create_sample_docx_bytes()
    res = client.post(
        "/api/documents",
        data={"specialty": "Neurology"},
        files={"file": ("neuro_report.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert res.status_code == 201
    doc_id = res.json()["id"]

    # Process
    pipeline = get_document_pipeline()
    pipeline.process_document(db=db_session, document_id=doc_id)

    # Verify
    detail_res = client.get(f"/api/documents/{doc_id}")
    assert detail_res.status_code == 200
    detail = detail_res.json()
    assert detail["status"] == DocumentProcessingStatus.COMPLETED.value
    assert detail["chunk_count"] >= 1


def test_strict_user_ownership_isolation(client: TestClient, db_session):
    # Register User A
    cookie_a = register_and_login_user(client, "owner_a@example.com", "Owner A")
    # Register User B
    cookie_b = register_and_login_user(client, "owner_b@example.com", "Owner B")

    # User A uploads a document
    client.cookies.set("session_id", cookie_a)
    res_a = client.post(
        "/api/documents",
        files={"file": ("confidential_a.txt", b"Confidential Medical Record A", "text/plain")},
    )
    assert res_a.status_code == 201
    doc_a_id = res_a.json()["id"]

    # User B switches session
    client.cookies.set("session_id", cookie_b)

    # 1. User B lists documents -> must be empty
    list_res = client.get("/api/documents")
    assert list_res.status_code == 200
    assert len(list_res.json()) == 0

    # 2. User B tries to view User A's document -> 404 Not Found
    get_res = client.get(f"/api/documents/{doc_a_id}")
    assert get_res.status_code == 404

    # 3. User B tries to delete User A's document -> 404 Not Found
    del_res = client.delete(f"/api/documents/{doc_a_id}")
    assert del_res.status_code == 404

    # User A can still view and delete their own document
    client.cookies.set("session_id", cookie_a)
    get_res_a = client.get(f"/api/documents/{doc_a_id}")
    assert get_res_a.status_code == 200

    del_res_a = client.delete(f"/api/documents/{doc_a_id}")
    assert del_res_a.status_code == 200

    # After deletion, User A also gets 404
    assert client.get(f"/api/documents/{doc_a_id}").status_code == 404


def test_upload_and_process_pdf_document(client: TestClient, db_session):
    cookie = register_and_login_user(client, "pdf_user@example.com", "PDF User")
    client.cookies.set("session_id", cookie)

    pdf_bytes = create_sample_pdf_bytes()
    res = client.post(
        "/api/documents",
        data={"specialty": "Cardiology"},
        files={"file": ("cardio_report.pdf", pdf_bytes, "application/pdf")},
    )
    assert res.status_code == 201
    doc_id = res.json()["id"]

    # Trigger pipeline
    pipeline = get_document_pipeline()
    pipeline.process_document(db=db_session, document_id=doc_id)

    # Verify document is in database with COMPLETED status
    detail_res = client.get(f"/api/documents/{doc_id}")
    assert detail_res.status_code == 200
    assert detail_res.json()["filename"] == "cardio_report.pdf"
    assert detail_res.json()["file_type"] == "pdf"
    assert detail_res.json()["status"] == DocumentProcessingStatus.COMPLETED.value
    assert detail_res.json()["chunk_count"] >= 1



def test_file_size_limit_rejection(client: TestClient, monkeypatch, db_session):
    cookie = register_and_login_user(client, "bigfile_user@example.com", "Big File")
    client.cookies.set("session_id", cookie)

    from app.core.config import settings
    monkeypatch.setattr(settings, "MAX_UPLOAD_SIZE_BYTES", 50)

    # Attempt to upload 100 bytes
    large_payload = b"A" * 100
    res = client.post(
        "/api/documents",
        files={"file": ("toolarge.txt", large_payload, "text/plain")},
    )
    assert res.status_code == 413
    assert "exceeds maximum allowed limit" in res.json()["detail"]


def test_processing_failure_status(client: TestClient, db_session):
    cookie = register_and_login_user(client, "fail_user@example.com", "Fail User")
    client.cookies.set("session_id", cookie)

    # Upload empty text content file that bypasses upload check (or fails during processing)
    # Uploading a txt file with only whitespace which TextCleaner reduces to empty
    res = client.post(
        "/api/documents",
        files={"file": ("whitespace_only.txt", b"   \n\n\t   ", "text/plain")},
    )
    assert res.status_code == 201
    doc_id = res.json()["id"]

    # Run processing directly and expect failure handled
    pipeline = get_document_pipeline()
    try:
        pipeline.process_document(db=db_session, document_id=doc_id)
    except Exception:
        pass

    # Verify document status is FAILED
    detail_res = client.get(f"/api/documents/{doc_id}")
    assert detail_res.status_code == 200
    detail = detail_res.json()
    assert detail["status"] == DocumentProcessingStatus.FAILED.value
    assert detail["processing_error"] is not None

